"""
build_excel_dashboard.py
------------------------
Generates a styled, multi-sheet Excel dashboard from the cleaned supply chain
data using openpyxl. Produces an executive Summary sheet (KPI cards), a
ChartData sheet with aggregated pivots, and native Excel charts.

Run (after data_cleaning.py):
    python src/build_excel_dashboard.py
Output:
    excel_analysis/Supplychain_Dashboard.xlsx
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.dataframe import dataframe_to_rows

ROOT = Path(__file__).resolve().parents[1]
CLEAN_PATH = ROOT / "data" / "Cleaned_Supplychain_Data.csv"
OUT_PATH = ROOT / "excel_analysis" / "Supplychain_Dashboard.xlsx"

NAVY = "1F3864"
BLUE = "2E5496"
LIGHT = "D9E1F2"
WHITE = "FFFFFF"

title_font = Font(name="Calibri", size=18, bold=True, color=WHITE)
kpi_label_font = Font(name="Calibri", size=10, bold=True, color=WHITE)
kpi_value_font = Font(name="Calibri", size=16, bold=True, color=WHITE)
header_font = Font(name="Calibri", size=11, bold=True, color=WHITE)
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _write_df(ws, df: pd.DataFrame, start_row: int = 1, start_col: int = 1):
    for r_off, row in enumerate(dataframe_to_rows(df, index=False, header=True)):
        for c_off, val in enumerate(row):
            cell = ws.cell(row=start_row + r_off, column=start_col + c_off, value=val)
            if r_off == 0:
                cell.font = header_font
                cell.fill = _fill(BLUE)
                cell.alignment = Alignment(horizontal="center")
            cell.border = border
    return start_row + len(df) + 1


def build_summary(ws, df: pd.DataFrame) -> None:
    ws.sheet_view.showGridLines = False
    ws.merge_cells("B2:H2")
    t = ws["B2"]
    t.value = "Supply Chain Performance Dashboard"
    t.font = title_font
    t.fill = _fill(NAVY)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 32

    kpis = [
        ("Total Revenue", f"${df['revenue'].sum()/1e6:,.1f}M"),
        ("Total Profit", f"${df['profit'].sum()/1e6:,.1f}M"),
        ("Avg Margin", f"{df['profit_margin'].mean()*100:,.1f}%"),
        ("Orders", f"{len(df):,}"),
        ("Avg Delivery", f"{df['delivery_days'].mean():,.1f} d"),
        ("On-Time (<=5d)", f"{df['on_time'].mean()*100:,.1f}%"),
    ]
    col = 2
    for label, value in kpis:
        lc = ws.cell(row=4, column=col, value=label)
        lc.font = kpi_label_font
        lc.fill = _fill(BLUE)
        lc.alignment = Alignment(horizontal="center")
        vc = ws.cell(row=5, column=col, value=value)
        vc.font = kpi_value_font
        vc.fill = _fill(NAVY)
        vc.alignment = Alignment(horizontal="center")
        ws.column_dimensions[vc.column_letter].width = 16
        col += 1


def build_chartdata(ws, df: pd.DataFrame) -> dict:
    """Write pivot tables to ChartData and return their cell ranges."""
    anchors = {}
    row = 1

    rev_channel = (
        df.groupby("Sales_Channel")["revenue"].sum().div(1e6).round(2).reset_index()
    )
    rev_channel.columns = ["Sales Channel", "Revenue ($M)"]
    anchors["rev_channel"] = (row, len(rev_channel))
    row = _write_df(ws, rev_channel, start_row=row) + 1

    wh = (
        df.groupby("WarehouseCode")["delivery_days"].mean().round(2).reset_index()
    )
    wh.columns = ["Warehouse", "Avg Delivery Days"]
    anchors["warehouse"] = (row, len(wh))
    row = _write_df(ws, wh, start_row=row) + 1

    monthly = (
        df.groupby("order_month")["revenue"].sum().div(1e6).round(2).reset_index()
    )
    monthly.columns = ["Month", "Revenue ($M)"]
    anchors["monthly"] = (row, len(monthly))
    row = _write_df(ws, monthly, start_row=row) + 1

    return anchors


def add_charts(wb, data_ws, anchors: dict, dash_ws) -> None:
    # Revenue by channel (pie)
    r0, n = anchors["rev_channel"]
    pie = PieChart()
    pie.title = "Revenue by Sales Channel"
    labels = Reference(data_ws, min_col=1, min_row=r0 + 1, max_row=r0 + n)
    data = Reference(data_ws, min_col=2, min_row=r0, max_row=r0 + n)
    pie.add_data(data, titles_from_data=True)
    pie.set_categories(labels)
    dash_ws.add_chart(pie, "B8")

    # Avg delivery days by warehouse (bar)
    r0, n = anchors["warehouse"]
    bar = BarChart()
    bar.type = "col"
    bar.title = "Avg Delivery Days by Warehouse"
    labels = Reference(data_ws, min_col=1, min_row=r0 + 1, max_row=r0 + n)
    data = Reference(data_ws, min_col=2, min_row=r0, max_row=r0 + n)
    bar.add_data(data, titles_from_data=True)
    bar.set_categories(labels)
    bar.legend = None
    dash_ws.add_chart(bar, "H8")

    # Monthly revenue trend (line)
    r0, n = anchors["monthly"]
    line = LineChart()
    line.title = "Monthly Revenue Trend ($M)"
    labels = Reference(data_ws, min_col=1, min_row=r0 + 1, max_row=r0 + n)
    data = Reference(data_ws, min_col=2, min_row=r0, max_row=r0 + n)
    line.add_data(data, titles_from_data=True)
    line.set_categories(labels)
    line.legend = None
    dash_ws.add_chart(line, "B24")


def main() -> None:
    df = pd.read_csv(CLEAN_PATH)
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    dash = wb.active
    dash.title = "Dashboard"
    build_summary(dash, df)

    data_ws = wb.create_sheet("ChartData")
    anchors = build_chartdata(data_ws, df)
    add_charts(wb, data_ws, anchors, dash)

    wb.save(OUT_PATH)
    print(f"Excel dashboard written to {OUT_PATH}")


if __name__ == "__main__":
    main()
